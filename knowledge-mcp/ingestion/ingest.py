import argparse
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, FieldCondition, Filter, MatchValue, PointStruct, VectorParams
from tqdm import tqdm
import yaml


PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR / "mcp"))
from policy_defaults import REQUIRED_DIRECTORY_NAMES, REQUIRED_TOP_LEVEL_NAMES, ensure_required_exclusions, remove_legacy_default_exclusions

load_dotenv(PROJECT_DIR / ".env")

BASE_DIR = Path(os.getenv("INGESTION_BASE_DIR", str(PROJECT_DIR / ".state")))
POLICY_FILE = Path(os.getenv("POLICY_FILE", str(PROJECT_DIR / "mcp" / "index-policy.yaml")))
STATE_FILE = BASE_DIR / "state" / "index-state.json"
ERROR_LOG = BASE_DIR / "logs" / "errors.log"
PROGRESS_FILE = Path(os.environ["INGEST_PROGRESS_FILE"]) if os.getenv("INGEST_PROGRESS_FILE") else None
_progress_last_write = 0
_progress_last_stage = None


def write_progress(stage, completed=0, total=None, force=False):
    global _progress_last_write, _progress_last_stage
    if PROGRESS_FILE is None:
        return
    now = time.monotonic()
    if not force and stage == _progress_last_stage and now - _progress_last_write < 0.25:
        return
    value = {"stage": stage, "completed": completed, "total": total, "updated_at": time.time()}
    temp = PROGRESS_FILE.with_suffix(".tmp")
    temp.write_text(json.dumps(value), encoding="utf-8")
    temp.replace(PROGRESS_FILE)
    _progress_last_write, _progress_last_stage = now, stage

COLLECTION_NAME = os.getenv("DENSE_COLLECTION", "documents")

MODEL_NAME = os.getenv("MODEL_NAME", "intfloat/multilingual-e5-small")

QDRANT_URL = os.getenv("QDRANT_URL", "http://127.0.0.1:6333")

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "1200"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "200"))
MIN_CHUNK_LENGTH = int(os.getenv("MIN_CHUNK_LENGTH", "80"))
EMBED_BATCH_SIZE = int(os.getenv("EMBED_BATCH_SIZE", "32"))


def load_policy():
    with POLICY_FILE.open("r", encoding="utf-8") as f:
        policy = yaml.safe_load(f)
    remove_legacy_default_exclusions(policy)
    ensure_required_exclusions(policy)
    root = os.getenv("KNOWLEDGE_ROOT") or policy.get("knowledge_root")
    if not root or root.startswith("${") or root == "/path/to/knowledge/root":
        raise ValueError("Set KNOWLEDGE_ROOT in .env or knowledge_root in index-policy.yaml")
    policy["knowledge_root"] = root
    return policy


def load_state():
    if not STATE_FILE.exists():
        return {"documents": {}}

    with STATE_FILE.open("r", encoding="utf-8") as f:
        return json.load(f)


def save_state(state):
    STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STATE_FILE.with_suffix(".tmp")

    with tmp.open("w", encoding="utf-8") as f:
        json.dump(state, f, indent=2, ensure_ascii=False)

    tmp.replace(STATE_FILE)


def log_error(path, error):
    ERROR_LOG.parent.mkdir(parents=True, exist_ok=True)

    with ERROR_LOG.open("a", encoding="utf-8") as f:
        timestamp = datetime.now(timezone.utc).isoformat()
        f.write(f"{timestamp}\t{path}\t{repr(error)}\n")


def sha256_file(path):
    digest = hashlib.sha256()

    with path.open("rb") as f:
        while True:
            block = f.read(1024 * 1024)

            if not block:
                break

            digest.update(block)

    return digest.hexdigest()


def is_candidate(path, root, policy):
    if not path.is_file():
        return False

    if not path.resolve().is_relative_to(root.resolve()):
        return False

    relative = path.relative_to(root)

    # Top-level exclusions
    if relative.parts:
        excluded_top = {name.casefold() for name in policy.get("exclude_top_level", [])} | REQUIRED_TOP_LEVEL_NAMES

        if relative.parts[0].casefold() in excluded_top:
            return False

    # Directory exclusions
    excluded_dirs = {name.casefold() for name in policy.get("exclude_directories", [])} | REQUIRED_DIRECTORY_NAMES

    if any(part.casefold() in excluded_dirs for part in relative.parts[:-1]):
        return False

    excluded_files = {name.casefold() for name in policy.get("exclude_files", [])}
    if path.name.casefold() in excluded_files or relative.as_posix().casefold() in excluded_files:
        return False

    # Extension policy
    extension = path.suffix.lower()

    included = set(policy.get("include_extensions", []))
    excluded = set(policy.get("exclude_extensions", []))

    if extension not in included:
        return False

    if extension in excluded:
        return False

    # File size
    max_size = policy.get("max_file_size_mb", 50) * 1024 * 1024

    try:
        size = path.stat().st_size
    except OSError:
        return False

    if size == 0 or size > max_size:
        return False

    return True


def discover_documents(policy):
    root = Path(policy["knowledge_root"]).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Knowledge root does not exist: {root}")

    candidates = []

    selected = json.loads(os.getenv("INDEX_SOURCE_PATHS", '[""]'))
    if not isinstance(selected, list) or not selected or not all(isinstance(item, str) for item in selected):
        raise ValueError("INDEX_SOURCE_PATHS must contain selected folders")
    for item in selected:
        if "\\" in item:
            raise ValueError("Invalid selected folder")
        folder = (root / item).resolve()
        if not folder.is_relative_to(root) or not folder.is_dir():
            raise ValueError(f"Selected folder is outside the document library: {item}")
        for path in folder.rglob("*"):
            if is_candidate(path, root, policy):
                candidates.append(path)

    return sorted(set(candidates))


def normalize_text(text):
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def parse_txt_or_md(path):
    text = path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    return [
        {
            "text": normalize_text(text),
            "page": None,
            "section": None,
        }
    ]


def parse_html(path):
    html = path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    soup = BeautifulSoup(html, "html.parser")

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    text = soup.get_text(separator="\n")

    return [
        {
            "text": normalize_text(text),
            "page": None,
            "section": None,
        }
    ]


def parse_pdf(path):
    reader = PdfReader(str(path))

    sections = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        text = normalize_text(text)

        if text:
            sections.append(
                {
                    "text": text,
                    "page": page_number,
                    "section": None,
                }
            )

    return sections


def parse_docx(path):
    document = Document(str(path))

    sections = []
    current_heading = None
    buffer = []

    def flush():
        nonlocal buffer

        text = normalize_text("\n".join(buffer))

        if text:
            sections.append(
                {
                    "text": text,
                    "page": None,
                    "section": current_heading,
                }
            )

        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        style = paragraph.style.name.lower() if paragraph.style else ""

        if style.startswith("heading"):
            flush()
            current_heading = text
        else:
            buffer.append(text)

    flush()

    return sections


def parse_pptx(path):
    presentation = Presentation(str(path))

    sections = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                value = shape.text.strip()

                if value:
                    parts.append(value)

        text = normalize_text("\n".join(parts))

        if text:
            sections.append(
                {
                    "text": text,
                    "page": slide_number,
                    "section": f"Slide {slide_number}",
                }
            )

    return sections


def parse_document(path):
    extension = path.suffix.lower()

    if extension in {".txt", ".md"}:
        return parse_txt_or_md(path)

    if extension in {".html", ".htm"}:
        return parse_html(path)

    if extension == ".pdf":
        return parse_pdf(path)

    if extension == ".docx":
        return parse_docx(path)

    if extension == ".pptx":
        return parse_pptx(path)

    raise ValueError(f"Unsupported extension: {extension}")


def chunk_text(text):
    if len(text) <= CHUNK_SIZE:
        if len(text) >= MIN_CHUNK_LENGTH:
            return [text]

        return []

    chunks = []

    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))

        chunk = text[start:end]

        # Prefer a natural break close to the end.
        if end < len(text):
            search_start = max(0, len(chunk) - 300)

            break_positions = [
                chunk.rfind("\n\n", search_start),
                chunk.rfind(". ", search_start),
                chunk.rfind("\n", search_start),
            ]

            best_break = max(break_positions)

            if best_break > 0:
                chunk = chunk[: best_break + 1]
                end = start + len(chunk)

        chunk = normalize_text(chunk)

        if len(chunk) >= MIN_CHUNK_LENGTH:
            chunks.append(chunk)

        if end >= len(text):
            break

        next_start = max(0, end - CHUNK_OVERLAP)

        # Safety against pathological loops
        if next_start <= start:
            next_start = end

        start = next_start

    return chunks


def make_point_id(document_hash, section_index, chunk_index):
    value = (
        f"{document_hash}:"
        f"{section_index}:"
        f"{chunk_index}"
    )

    # Qdrant accepts UUID or unsigned integer IDs.
    return int(
        hashlib.sha256(value.encode()).hexdigest()[:16],
        16,
    )


def ensure_collection(client, vector_size):
    existing = {
        collection.name
        for collection in client.get_collections().collections
    }

    if COLLECTION_NAME not in existing:
        client.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=VectorParams(
                size=vector_size,
                distance=Distance.COSINE,
            ),
        )
        return True
    actual_size = client.get_collection(COLLECTION_NAME).config.params.vectors.size
    if actual_size != vector_size:
        raise ValueError(f"Collection vector size {actual_size} differs from model size {vector_size}")
    return False


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse and chunk without embeddings/Qdrant writes.",
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Process only the first N unique documents.",
    )
    parser.add_argument(
        "--allow-empty",
        action="store_true",
        help="Allow an empty source to remove all indexed documents.",
    )

    args = parser.parse_args()
    if args.limit is not None and args.limit < 1:
        parser.error("--limit must be greater than zero")
    if not 0 <= CHUNK_OVERLAP < CHUNK_SIZE:
        parser.error("CHUNK_OVERLAP must be smaller than CHUNK_SIZE")

    start_time = time.time()

    policy = load_policy()
    root = Path(policy["knowledge_root"]).resolve()

    print("Discovering candidate documents...")
    write_progress("discovering")

    candidates = discover_documents(policy)

    print(f"Candidates found: {len(candidates)}")
    if not candidates and not args.dry_run and not args.allow_empty:
        raise SystemExit("No eligible documents found; the index was not changed. Use --allow-empty to clear it intentionally.")

    # -------------------------------------------------------------
    # Deduplication
    # -------------------------------------------------------------

    hash_to_paths = {}

    print("Calculating SHA-256 hashes...")

    hash_failures = 0
    for count, path in enumerate(tqdm(candidates), 1):
        try:
            digest = sha256_file(path)

            hash_to_paths.setdefault(
                digest,
                [],
            ).append(path)

        except Exception as exc:
            hash_failures += 1
            log_error(path, exc)
        write_progress("hashing", count, len(candidates))
    write_progress("hashing", len(candidates), len(candidates), force=True)

    unique_documents = [
        (digest, paths)
        for digest, paths in hash_to_paths.items()
    ]

    print(f"Unique documents: {len(unique_documents)}")
    print(
        "Duplicate files skipped:",
        len(candidates) - len(unique_documents),
    )

    if args.limit is not None:
        unique_documents = unique_documents[: args.limit]

    # -------------------------------------------------------------
    # Dry-run parser/chunk test
    # -------------------------------------------------------------

    if args.dry_run:
        if not unique_documents:
            raise SystemExit("No eligible documents found for the dry run.")
        print("\nDRY RUN - no embeddings or Qdrant writes\n")

        total_chunks = 0
        parsed = 0
        failed = 0

        for count, (digest, paths) in enumerate(tqdm(unique_documents), 1):
            primary = paths[0]

            try:
                sections = parse_document(primary)

                document_chunks = 0

                for section in sections:
                    document_chunks += len(
                        chunk_text(section["text"])
                    )

                parsed += 1
                total_chunks += document_chunks

            except Exception as exc:
                failed += 1
                log_error(primary, exc)
            write_progress("dry_run", count, len(unique_documents))
        write_progress("dry_run", len(unique_documents), len(unique_documents), force=True)

        print("\n=== DRY RUN RESULTS ===")
        print(f"Documents parsed: {parsed}")
        print(f"Failed:           {failed}")
        print(f"Chunks produced:  {total_chunks}")

        elapsed = time.time() - start_time

        print(f"Elapsed:          {elapsed:.1f}s")

        if failed or hash_failures:
            raise SystemExit(1)
        write_progress("complete", parsed, parsed, force=True)
        return

    # -------------------------------------------------------------
    # Full ingestion
    # -------------------------------------------------------------

    print("\nLoading embedding model:")
    print(MODEL_NAME)
    write_progress("loading_model")

    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(
        MODEL_NAME,
        device="cpu",
    )

    client = QdrantClient(url=QDRANT_URL)

    collection_created = ensure_collection(client, model.get_sentence_embedding_dimension())

    state = load_state()
    if collection_created or client.count(COLLECTION_NAME, exact=True).count == 0:
        state = {"documents": {}}

    indexed_documents = 0
    skipped_documents = 0
    failed_documents = 0
    indexed_chunks = 0

    write_progress("indexing", 0, len(unique_documents), force=True)
    for count, (digest, paths) in enumerate(tqdm(unique_documents), 1):
        write_progress("indexing", count - 1, len(unique_documents))
        primary = paths[0]

        relative_primary = primary.relative_to(root).as_posix()

        duplicate_sources = [
            path.relative_to(root).as_posix()
            for path in paths[1:]
        ]

        # Already indexed and unchanged.
        if state["documents"].get(digest, {}).get("source") == relative_primary and state["documents"][digest].get("duplicate_sources") == duplicate_sources:
            skipped_documents += 1
            continue

        try:
            sections = parse_document(primary)

            chunk_records = []

            for section_index, section in enumerate(sections):
                chunks = chunk_text(section["text"])

                for chunk_index, text in enumerate(chunks):
                    chunk_records.append(
                        {
                            "section_index": section_index,
                            "chunk_index": chunk_index,
                            "text": text,
                            "page": section["page"],
                            "section": section["section"],
                        }
                    )

            if not chunk_records:
                skipped_documents += 1
                continue

            # E5 models work best with explicit passage/query prefixes.
            embedding_inputs = [
                f"passage: {record['text']}"
                for record in chunk_records
            ]

            embeddings = model.encode(
                embedding_inputs,
                batch_size=EMBED_BATCH_SIZE,
                normalize_embeddings=True,
                show_progress_bar=False,
            )

            stat = primary.stat()

            points = []

            for record, vector in zip(
                chunk_records,
                embeddings,
            ):
                point_id = make_point_id(
                    digest,
                    record["section_index"],
                    record["chunk_index"],
                )

                payload = {
                    "document_id": digest,
                    "file_hash": digest,

                    "source": relative_primary,
                    "duplicate_sources": duplicate_sources,

                    "filename": primary.name,
                    "extension": primary.suffix.lower(),

                    "section_index": record["section_index"],
                    "chunk_index": record["chunk_index"],

                    "page": record["page"],
                    "section": record["section"],

                    "text": record["text"],

                    "modified_at": stat.st_mtime,
                    "file_size": stat.st_size,
                }

                points.append(
                    PointStruct(
                        id=point_id,
                        vector=vector.tolist(),
                        payload=payload,
                    )
                )

            # Avoid very large single requests.
            for offset in range(0, len(points), 100):
                client.upsert(
                    collection_name=COLLECTION_NAME,
                    points=points[offset : offset + 100],
                    wait=True,
                )

            state["documents"][digest] = {
                "source": relative_primary,
                "duplicate_sources": duplicate_sources,
                "chunks": len(points),
                "indexed_at": datetime.now(
                    timezone.utc
                ).isoformat(),
            }

            save_state(state)

            indexed_documents += 1
            indexed_chunks += len(points)

        except Exception as exc:
            failed_documents += 1
            log_error(primary, exc)

    write_progress("indexing", len(unique_documents), len(unique_documents), force=True)

    elapsed = time.time() - start_time

    if args.limit is None and not failed_documents and not hash_failures:
        for digest in set(state["documents"]) - set(hash_to_paths):
            client.delete(
                collection_name=COLLECTION_NAME,
                points_selector=Filter(must=[FieldCondition(key="document_id", match=MatchValue(value=digest))]),
                wait=True,
            )
            del state["documents"][digest]
        save_state(state)

    print("\n=== INGESTION COMPLETE ===")
    print(f"Indexed documents: {indexed_documents}")
    print(f"Already indexed:   {skipped_documents}")
    print(f"Failed:            {failed_documents}")
    print(f"Chunks indexed:    {indexed_chunks}")
    print(f"Elapsed:           {elapsed / 60:.1f} min")
    if failed_documents or hash_failures:
        raise SystemExit(1)
    write_progress("complete", len(unique_documents), len(unique_documents), force=True)


if __name__ == "__main__":
    main()
